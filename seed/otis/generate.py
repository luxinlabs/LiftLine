"""Build the Otis hackathon demo database.

Creates otis_hackathon.db (SQLite) from schema.sql, seeds deterministic sample
data (background noise + the scripted demo scenarios), then exports every table
to csv/ and to one Excel workbook.

    python3 generate.py

Demo "now" is 2026-09-26 10:15 site-local time.
"""
import csv
import os
import random
import sqlite3
from datetime import date, datetime, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(HERE, "otis_hackathon.db")
NOW = datetime(2026, 9, 26, 10, 15)
TODAY = NOW.date()
rng = random.Random(42)

rows = {t: [] for t in [
    "customers", "sites", "equipment", "service_contracts", "parts", "warranties",
    "contacts", "error_code_catalog", "error_events", "technicians",
    "service_cases", "work_orders", "work_order_parts", "inspections", "access_credentials"]}


def ts(d):
    return d.strftime("%Y-%m-%d %H:%M:%S") if d else None


def ds(d):
    return d.isoformat() if d else None


# ------------------------------------------------------------------ reference data
ERROR_CODES = [
    # code, category, description, severity, safety, causes, action, customer text, threshold
    ("DLC-101", "Door Lock / Safety Chain", "Landing door lock circuit interrupted during travel", "High", 1,
     "Worn or misaligned landing door lock contact; loose door hanger; debris in sill",
     "Inspect lock contact and roller alignment at reported floor; check door hanger and sill",
     "A safety sensor on one of the landing doors briefly signalled 'not locked', so the elevator stopped as a precaution.", 10),
    ("DLC-102", "Door Lock / Safety Chain", "Car door contact open during travel", "High", 1,
     "Car door contact wear; car door operator clutch misadjusted",
     "Inspect car door contact and clutch engagement", "A door sensor on the car itself interrupted the trip for safety.", 8),
    ("REL-210", "Leveling", "Automatic re-level to nearest landing and door cycle after safety stop", "Info", 0,
     "Normal recovery behaviour after a safety chain interruption", "No action on its own; correlate with preceding fault",
     "The elevator moved to the nearest floor and opened its doors so passengers could exit safely.", None),
    ("LEV-220", "Leveling", "Leveling error exceeds tolerance (>10 mm)", "Medium", 0,
     "Encoder drift; brake wear; rope stretch", "Check leveling sensors and brake", "The car stopped slightly above or below the floor.", 15),
    ("LEV-221", "Leveling", "Repeated re-leveling at landing", "Low", 0,
     "Hydraulic valve leak-down; load variation", "Check valve and oil level (hydraulic units)", "The car adjusted its position at the floor.", 25),
    ("DPD-305", "Door Protection", "Door reversal by light curtain / protective device", "Info", 0,
     "Obstruction in doorway; passengers holding doors; dirty light curtain lens",
     "If very frequent, clean/align light curtain; advise building on obstructions",
     "The doors reopened because something was in the doorway.", None),
    ("DPD-306", "Door Protection", "Door nudging mode activated (door held open too long)", "Low", 0,
     "Doors held open; persistent obstruction", "Advise building; check light curtain", "The doors closed slowly because they had been held open.", None),
    ("DOP-310", "Door Operator", "Door operator motor overcurrent", "Medium", 0,
     "Door track debris; worn belt; operator motor wear", "Clean track, inspect belt and operator", "The door motor worked harder than normal.", 10),
    ("DOP-311", "Door Operator", "Door close timeout", "Medium", 0,
     "Obstruction; misaligned door panels", "Inspect door panels and guides", "The doors took too long to close.", 12),
    ("DRV-401", "Drive", "Drive overcurrent trip", "High", 0,
     "Motor or drive fault; overload", "Inspect drive and motor; review load data", "The elevator's motor drive shut down to protect itself.", 3),
    ("DRV-402", "Drive", "Drive overtemperature", "Medium", 0,
     "Machine room ventilation; heavy duty cycle", "Check machine room cooling and drive fans", "The motor drive ran hot.", 5),
    ("BRK-450", "Brake", "Brake not lifting within time", "High", 1,
     "Brake coil or lining wear", "Inspect brake immediately", "A brake check did not complete as expected.", 2),
    ("SAF-501", "Safety Chain", "Emergency stop - safety chain open (unspecified)", "High", 1,
     "Any safety device in the chain", "Check safety chain status log", "A safety device stopped the elevator.", 3),
    ("SAF-510", "Safety Chain", "Overspeed governor switch tripped", "Critical", 1,
     "Governor fault or overspeed", "Take unit out of service; dispatch immediately", "The elevator's overspeed safety system activated.", 1),
    ("ENT-900", "Entrapment", "Passenger entrapment detected (car stopped with occupancy, doors closed > 90 s)", "Critical", 1,
     "Any stop between floors with passengers", "P1: dispatch immediately, stay on line with passengers",
     "Passengers may be inside a stopped elevator.", 1),
    ("ALM-910", "Alarm", "Car alarm button pressed", "Medium", 1,
     "Passenger alarm; possible entrapment or misuse", "Voice-connect to car; verify", "Someone pressed the alarm button in the elevator.", None),
    ("PWR-601", "Power", "Mains power loss", "Medium", 0,
     "Building power interruption", "Confirm with building; verify rescue operation", "The elevator lost building power.", None),
    ("COM-701", "Communication", "Remote monitoring heartbeat lost", "Low", 0,
     "Network outage at site; gateway fault", "Check site connectivity", "We temporarily lost the monitoring connection.", None),
    ("HYD-801", "Hydraulic", "Low oil pressure", "Medium", 0,
     "Leak; pump wear", "Inspect hydraulic system", "Hydraulic pressure was lower than expected.", 5),
    ("SEC-960", "Floor Security", "Car call to secured floor rejected - badge not enabled for floor in elevator security table", "Low", 0,
     "Badge not authorized; floor-security mapping out of sync with access control system",
     "Compare badge rights in access control with elevator floor-security table; reprogram if mismatched",
     "The elevator did not accept the badge for that floor, so it would not take you there.", 5),
    ("SEC-961", "Floor Security", "Car call to secured floor accepted (badge validated)", "Info", 0,
     "Normal operation", "None", "Your badge was accepted for that floor.", None),
    ("CAL-950", "Call System", "Hall call button stuck", "Low", 0,
     "Button wear; moisture", "Replace hall button", "A call button on a floor appears stuck.", 20),
]

PARTS = [
    ("DL-CNT-220", "Landing door lock contact assembly", "Door Lock", 185.0, 90, "OTC-400,OTC-600"),
    ("DL-RLR-110", "Landing door lock roller set", "Door Lock", 64.0, 90, "OTC-400,OTC-600,OTC-200"),
    ("DH-ASM-300", "Landing door hanger assembly", "Door Hardware", 420.0, 180, "ALL"),
    ("CD-CNT-120", "Car door contact", "Door Lock", 140.0, 90, "ALL"),
    ("DO-BLT-050", "Door operator drive belt", "Door Operator", 58.0, 90, "ALL"),
    ("DO-MTR-700", "Door operator motor", "Door Operator", 1250.0, 365, "OTC-400,OTC-600"),
    ("LC-CRT-900", "Light curtain pair", "Door Protection", 980.0, 365, "ALL"),
    ("EN-SNS-410", "Leveling sensor", "Leveling", 310.0, 180, "ALL"),
    ("BR-LIN-800", "Brake lining kit", "Brake", 760.0, 365, "ALL"),
    ("DR-FAN-120", "Drive cooling fan", "Drive", 190.0, 180, "ALL"),
    ("HB-BTN-010", "Hall call button", "Fixtures", 45.0, 90, "ALL"),
    ("HY-VLV-500", "Hydraulic control valve", "Hydraulic", 1850.0, 365, "OTC-200"),
]

TECHS = [
    # id, name, region, city, certs, level, status, lat, lon
    ("T-501", "Mike Johnson", "SF Bay - North", "San Francisco", "OTC-400,OTC-600", "Senior", "Available", 37.7880, -122.4075),
    ("T-502", "Megan Doyle", "SF Bay - North", "Oakland", "OTC-200,OTC-400", "Mechanic", "Available", 37.8044, -122.2712),
    ("T-503", "Luke Harris", "SF Bay - North", "San Francisco", "OTC-600", "Mechanic", "On Job", 37.7599, -122.4148),
    ("T-504", "Adam King", "SF Bay - South", "San Jose", "OTC-400,OTC-600", "Specialist", "Available", 37.3382, -121.8863),
    ("T-505", "Grace Lewis", "SF Bay - South", "Palo Alto", "OTC-200,OTC-400", "Senior", "Off Shift", 37.4419, -122.1430),
    ("T-506", "James Carter", "SF Bay - East", "Walnut Creek", "OTC-200,OTC-400,OTC-600", "Senior", "Available", 37.9101, -122.0652),
    ("T-507", "Sophie Ross", "SF Bay - East", "Fremont", "OTC-400", "Apprentice", "Available", 37.5485, -121.9886),
    ("T-508", "Ben Walker", "SF Bay - South", "Santa Clara", "OTC-600", "Mechanic", "On Job", 37.3541, -121.9552),
    ("T-509", "Hannah White", "SF Bay - North", "Daly City", "OTC-200", "Mechanic", "Available", 37.6879, -122.4702),
    ("T-510", "David Moore", "SF Bay - East", "Hayward", "OTC-200,OTC-600", "Mechanic", "Off Shift", 37.6688, -122.0808),
]
REGION_CITIES = {
    "SF Bay - North": [("San Francisco", 37.78, -122.41), ("Daly City", 37.69, -122.47), ("Oakland", 37.80, -122.27)],
    "SF Bay - South": [("San Jose", 37.34, -121.89), ("Palo Alto", 37.44, -122.14), ("Santa Clara", 37.35, -121.96)],
    "SF Bay - East": [("Walnut Creek", 37.91, -122.07), ("Fremont", 37.55, -121.99), ("Hayward", 37.67, -122.08)],
}
FIRST = ["James", "Mary", "John", "Emma", "Peter", "Anna", "Paul", "Lucy", "Mark", "Kate", "David", "Amy",
         "Chris", "Laura", "Tom", "Rachel", "Steve", "Julia", "Sam", "Helen", "Nick", "Jane", "Rob", "Sue"]
LAST = ["Smith", "Brown", "Taylor", "Wilson", "Davies", "Evans", "Thomas", "Roberts", "Walker", "Wright",
        "Green", "Hall", "Wood", "Baker", "Hill", "Cooper", "Ward", "Turner", "Parker", "Bell"]
STREETS = ["Market St", "Mission St", "Howard St", "Broadway", "El Camino Real", "Main St", "Park Ave",
           "University Ave", "Oak St", "Lake St", "Harbor Blvd", "Central Ave"]

# ------------------------------------------------------------------ counters
seq = {"cust": 1000, "site": 2000, "eq": 30000, "con": 5000, "ct": 8000, "wty": 9000,
       "wo": 700000, "case": 47000, "evt": 0, "insp": 6000}


def nxt(key):
    seq[key] += 1
    return seq[key]


def phone():
    return "+1-415-%03d-%04d" % (rng.randint(200, 999), rng.randint(0, 9999))


# ------------------------------------------------------------------ builders
def add_customer(name, segment, tier, city, cid=None):
    n = nxt("cust") if cid is None else cid
    customer_id = "CU-%d" % n
    rows["customers"].append(dict(
        customer_id=customer_id, sap_customer_no="10%06d" % n, crm_account_id="001Hx0000%06dAAA" % n,
        name=name, segment=segment, account_tier=tier, billing_city=city, billing_state="CA", country="US",
        created_at=ds(date(rng.randint(2008, 2022), rng.randint(1, 12), 1))))
    return customer_id


def add_site(customer_id, name, city, region, floors, btype, lat, lon, address=None, postal=None):
    n = nxt("site")
    site_id = "ST-%d" % n
    rows["sites"].append(dict(
        site_id=site_id, customer_id=customer_id, sap_functional_location="US-CA-%s-%d" % (region[-5:].strip().upper()[:3], n),
        name=name, address=address or "%d %s" % (rng.randint(10, 2500), rng.choice(STREETS)), city=city, state="CA",
        postal_code=postal or "9%04d" % rng.randint(4000, 5999), building_type=btype, floors=floors, service_region=region,
        timezone="America/Los_Angeles", latitude=round(lat, 4), longitude=round(lon, 4)))
    return site_id


def add_equipment(site_id, label, product_line, controller, floors, install, drive="Gearless traction", status="In Service",
                  notes=None, secured=None):
    n = nxt("eq")
    equipment_id = "EQ-%d" % n
    rows["equipment"].append(dict(
        equipment_id=equipment_id, sap_equipment_no="E%09d" % n, site_id=site_id, car_label=label,
        equipment_type="Elevator", product_line=product_line, controller_model=controller, drive_type=drive,
        floors_served=floors, capacity_kg=rng.choice([1000, 1150, 1350, 1600]),
        speed_mps=rng.choice([1.0, 1.6, 2.0, 2.5]) if drive != "Hydraulic" else 0.6,
        install_date=ds(install), iot_device_id="REM-%07d" % (n * 7 % 9999991),
        remote_monitoring_enabled=1, operational_status=status, lowest_floor_served="L",
        highest_floor_served=str(floors), service_notes=notes, secured_floors=secured))
    # unit warranty: 24 months from install
    add_warranty(equipment_id, "Unit", None, None, "Complete unit (new installation)", install,
                 install + timedelta(days=730), "24-month new-installation warranty: parts and labor")
    return equipment_id


def add_warranty(equipment_id, level, part_no, wo_id, component, start, end, terms):
    rows["warranties"].append(dict(
        warranty_id="WT-%d" % nxt("wty"), equipment_id=equipment_id, coverage_level=level, part_no=part_no,
        work_order_id=wo_id, component=component, start_date=ds(start), end_date=ds(end), terms=terms))


def add_contract(customer_id, site_id, tier, start, end):
    n = nxt("ct")
    comp = tier == "Comprehensive"
    rows["service_contracts"].append(dict(
        contract_id="CT-%d" % n, sap_contract_no="40%08d" % n, customer_id=customer_id, site_id=site_id, tier=tier,
        start_date=ds(start), end_date=ds(end), status="Active" if end >= TODAY else "Expired",
        parts_covered=1 if tier != "Basic" else 0, labor_covered=1, after_hours_covered=1 if comp else 0,
        response_sla_hours={"Comprehensive": 2, "Standard": 4, "Basic": 8}[tier],
        repeat_fault_sla_hours={"Comprehensive": 4, "Standard": 8, "Basic": 24}[tier],
        callback_window_days={"Comprehensive": 30, "Standard": 14, "Basic": 7}[tier],
        non_contract_callout_usd=395.0, annual_value_usd=float(rng.randint(6, 40) * 1000)))


def add_contact(customer_id, site_id, first, last, role, authorized=1, status="Active", ph=None, channel="Phone",
                company=None):
    n = nxt("con")
    contact_id = "CN-%d" % n
    rows["contacts"].append(dict(
        contact_id=contact_id, customer_id=customer_id, site_id=site_id, first_name=first, last_name=last, role=role, company=company,
        phone=ph or phone(), email="%s.%s@example.com" % (first.lower(), last.lower()),
        is_authorized_for_service=authorized, preferred_channel=channel, status=status))
    return contact_id


def add_event(equipment_id, code, when, floor, position="At Landing", direction="Idle", door="Closed", related=None,
              credential=None):
    seq["evt"] += 1
    event_id = "EV-%08d" % seq["evt"]
    eq = next(e for e in rows["equipment"] if e["equipment_id"] == equipment_id)
    rows["error_events"].append(dict(
        event_id=event_id, equipment_id=equipment_id, iot_device_id=eq["iot_device_id"], error_code=code,
        event_ts=ts(when), floor=floor, car_position=position, direction=direction, load_pct=rng.randint(0, 85),
        door_state=door, related_event_id=related, credential_id=credential, acknowledged=1 if when < NOW - timedelta(days=3) else 0))
    return event_id


def add_case(customer_id, site_id, equipment_id, contact_id, opened, channel, category, subject, desc, priority,
             status, resolved=None, resolution=None, safety=0, repeat=0, related=None, sla_hours=4):
    n = nxt("case")
    case_id = "CS-%d" % n
    rows["service_cases"].append(dict(
        case_id=case_id, crm_case_id="500Hx0000%07dQAA" % n, customer_id=customer_id, site_id=site_id,
        equipment_id=equipment_id, contact_id=contact_id, opened_at=ts(opened), channel=channel, category=category,
        subject=subject, description=desc, priority=priority, safety_flag=safety, is_repeat_fault=repeat,
        related_case_id=related, status=status, sla_due_at=ts(opened + timedelta(hours=sla_hours)),
        resolved_at=ts(resolved), resolution_summary=resolution))
    return case_id


def add_wo(equipment_id, wtype, tech, scheduled, floor=None, component=None, replaced=0, findings=None, action=None,
           case_id=None, parent=None, billable=0, status="Completed", hours=None, parts=()):
    n = nxt("wo")
    wo_id = "WO-%d" % n
    done = status == "Completed"
    started = scheduled + timedelta(minutes=rng.randint(0, 40)) if done else None
    hrs = hours or round(rng.uniform(0.75, 3.5), 2)
    rows["work_orders"].append(dict(
        work_order_id=wo_id, fsm_work_order_no="SA-%08d" % n, case_id=case_id, equipment_id=equipment_id,
        parent_work_order_id=parent, work_type=wtype, technician_id=tech, scheduled_at=ts(scheduled),
        started_at=ts(started), completed_at=ts(started + timedelta(hours=hrs)) if done else None, floor=floor,
        component=component, component_replaced=replaced, findings=findings, action_taken=action,
        labor_hours=hrs if done else None, billable=billable, status=status))
    for part_no, qty in parts:
        rows["work_order_parts"].append(dict(work_order_id=wo_id, part_no=part_no, quantity=qty))
        days = next(p[4] for p in PARTS if p[0] == part_no)
        comp = next(p[1] for p in PARTS if p[0] == part_no)
        end_day = (started or scheduled).date()
        add_warranty(equipment_id, "Component", part_no, wo_id, comp + (" - floor %d" % floor if floor else ""),
                     end_day, end_day + timedelta(days=days), "%d-day parts warranty on replaced component" % days)
    return wo_id


def pm_history(equipment_id, techs, months=12):
    """Monthly preventive maintenance visits."""
    for m in range(months, 0, -1):
        d = datetime.combine(TODAY - timedelta(days=30 * m + rng.randint(-5, 5)), datetime.min.time()).replace(
            hour=rng.randint(8, 14))
        add_wo(equipment_id, "Preventive Maintenance", rng.choice(techs), d, component="Routine maintenance",
               findings="Routine checks completed; no defects found" if rng.random() > 0.2 else
               "Minor wear noted on door gibs; lubricated rails and door tracks",
               action="Performed monthly maintenance checklist", hours=round(rng.uniform(0.75, 1.5), 2))


def inspection(equipment_id, last_date, result="Pass", notes=None):
    n = nxt("insp")
    rows["inspections"].append(dict(
        inspection_id="IN-%d" % n, equipment_id=equipment_id, inspection_date=ds(last_date),
        authority="California DIR - Elevator Unit", result=result, notes=notes,
        certificate_no="CA-EL-%06d" % rng.randint(100000, 999999), next_due_date=ds(last_date + timedelta(days=365))))


# ================================================================== SCENARIOS
# ---- Main story: Harbor Point Tower, Car B, repeat door-lock fault ----------
crest = add_customer("Crestline Properties", "Commercial Office", "Key", "San Francisco")
harbor = add_site(crest, "Harbor Point Tower", "San Francisco", "SF Bay - North", 18, "Office", 37.7913, -122.3931,
                  address="201 Spear St", postal="94105")
ROOF = "Serves L-18. Floor 18 is the rooftop terrace: secured floor, badge must be tapped on the in-car reader"
car_a = add_equipment(harbor, "Car A", "Gen2", "OTC-600", 18, date(2019, 4, 15), notes=ROOF, secured="18")
car_b = add_equipment(harbor, "Car B", "Gen2", "OTC-600", 18, date(2019, 4, 15), notes=ROOF, secured="18")
car_c = add_equipment(harbor, "Car C", "Gen2", "OTC-600", 18, date(2019, 4, 15), notes=ROOF, secured="18")
add_contract(crest, harbor, "Comprehensive", date(2024, 1, 1), date(2026, 12, 31))
sarah = add_contact(crest, harbor, "Sarah", "Miller", "Facilities Manager", ph="+1-415-555-0142", channel="SMS")
add_contact(crest, harbor, "Tom", "Baker", "Former Facilities Manager", authorized=1, status="Inactive",
            ph="+1-415-555-0199")
add_contact(crest, None, "Linda", "Moss", "VP Property Operations", ph="+1-415-555-0110", channel="Email")
for car in (car_a, car_b, car_c):
    pm_history(car, ["T-501", "T-503"])
    inspection(car, date(2026, 3, 18))

# the first occurrence, 3 weeks ago, reported by Sarah
first_fault = datetime(2026, 9, 4, 16, 40)
for i in range(3):
    t = first_fault + timedelta(hours=i * 7 + rng.randint(0, 3))
    ev = add_event(car_b, "DLC-101", t, 7, "Between Floors", rng.choice(["Up", "Down"]))
    add_event(car_b, "REL-210", t + timedelta(seconds=25), 7, "At Landing", "Idle", "Open")
prev_case = add_case(
    crest, harbor, car_b, sarah, datetime(2026, 9, 5, 9, 5), "Phone", "Breakdown",
    "Car B stopping between floors near 7",
    "Caller reports Car B stops suddenly, moves to a floor, opens and closes doors, then continues. Happened 3 times yesterday.",
    "P2", "Resolved", resolved=datetime(2026, 9, 7, 13, 30),
    resolution="Replaced landing door lock contact on floor 7 (Car B). Tested 20 runs, no recurrence.", sla_hours=2)
prev_wo = add_wo(
    car_b, "Repair", "T-501", datetime(2026, 9, 7, 10, 0), floor=7, component="Landing door lock contact",
    replaced=1, case_id=prev_case,
    findings="Landing door lock contact at floor 7 worn and intermittently opening safety chain; roller slightly misaligned",
    action="Replaced landing door lock contact assembly (DL-CNT-220); adjusted roller alignment; test-ran 20 trips",
    hours=2.5, parts=[("DL-CNT-220", 1)])

# the recurrence: 14 events over the last 6 days (11 at floor 7), nobody has called until now
starts = sorted(NOW - timedelta(days=6) + timedelta(minutes=rng.randint(0, 6 * 24 * 60 - 120)) for _ in range(14))
floors = [7] * 11 + [6, 8, 8]
rng.shuffle(floors)
for t, fl in zip(starts, floors):
    t = t.replace(hour=rng.randint(7, 19))
    rel = add_event(car_b, "REL-210", t + timedelta(seconds=rng.randint(15, 40)), fl, "At Landing", "Idle", "Open")
    add_event(car_b, "DLC-101", t, fl, "Between Floors", rng.choice(["Up", "Down"]), "Closed", related=rel)
# cars A and C: only normal door reversals
for car in (car_a, car_c):
    for _ in range(rng.randint(3, 6)):
        add_event(car, "DPD-305", NOW - timedelta(minutes=rng.randint(60, 20000)), rng.randint(1, 17), door="Opening")

# ---- Rooftop story: CodeRabbit staff can't get to floor 18 ------------------
emily = add_contact(crest, harbor, "Emily", "Clark", "Office Manager", company="CodeRabbit", ph="+1-415-555-0171")
jack = add_contact(crest, harbor, "Jack", "Wilson", "Head of Workplace", company="CodeRabbit", ph="+1-415-555-0172")
creds = [
    ("AC-10001", emily, "Emily Clark", "CodeRabbit", "L,15,16,18"),
    ("AC-10002", jack, "Jack Wilson", "CodeRabbit", "L,15,16,18"),
    ("AC-10003", sarah, "Sarah Miller", "Crestline Properties", "ALL"),
    ("AC-10004", None, "Paul Wright", "Harbor Legal LLP", "L,9,10,18"),
    ("AC-10005", None, "Anna Cooper", "Harbor Legal LLP", "L,9,10"),
    ("AC-10006", None, "Nick Ward", "Bayside Design Co", "L,12,18"),
]
for cid, contact, holder, company, floors_ in creds:
    rows["access_credentials"].append(dict(
        credential_id=cid, site_id=harbor, contact_id=contact, holder_name=holder, company=company,
        badge_no="B%05d" % rng.randint(10000, 99999), authorized_floors=floors_, status="Active",
        updated_at=ts(datetime(2026, rng.randint(1, 8), rng.randint(1, 28), 9, 0))))
# 23 Sep: building asked for new tenant floors 14-16 to be added to the floor-security table
reprog = datetime(2026, 9, 23, 11, 0)
add_wo(car_a, "Programming", "T-503", reprog, floor=18, component="Floor-security table (all cars)",
       findings="Building security requested badge access for new tenant floors 14-16",
       action="Updated floor-security mapping on Cars A, B, C for floors 14-16; tested floors 14-16 OK",
       hours=1.5)
roof_users = ["AC-10001", "AC-10002", "AC-10004", "AC-10006"]
# before the change: floor 18 calls accepted
for _ in range(25):
    t = reprog - timedelta(days=rng.randint(1, 20), hours=rng.randint(0, 8), minutes=rng.randint(0, 59))
    add_event(rng.choice([car_a, car_b, car_c]), "SEC-961", t.replace(hour=rng.randint(10, 18)), 18,
              direction="Up", credential=rng.choice(roof_users))
# after the change: every floor 18 call rejected, including the two CodeRabbit badges
for _ in range(19):
    t = (reprog + timedelta(days=rng.randint(0, 3))).replace(hour=rng.randint(13, 19), minute=rng.randint(0, 59))
    if t >= NOW or t < reprog + timedelta(hours=3):
        continue
    add_event(rng.choice([car_a, car_b, car_c]), "SEC-960", t, 18, direction="Up",
              credential=rng.choice(roof_users + ["AC-10001", "AC-10002"]))
# floors 14-16 working fine after the change
for _ in range(12):
    t = (reprog + timedelta(days=rng.randint(0, 3))).replace(hour=rng.randint(8, 18), minute=rng.randint(0, 59))
    if reprog + timedelta(hours=2) < t < NOW:
        add_event(rng.choice([car_a, car_b, car_c]), "SEC-961", t, rng.choice([15, 16]), direction="Up",
                  credential=rng.choice(["AC-10001", "AC-10002"]))

# ---- Alt 2: Maple Court - no contract, warranty expired --------------------
maple_c = add_customer("Maple Court Homeowners Association", "Residential", "Standard", "Oakland")
maple = add_site(maple_c, "Maple Court Apartments", "Oakland", "SF Bay - North", 6, "Residential", 37.8120, -122.2600)
maple_eq = add_equipment(maple, "Main Elevator", "Gen2", "OTC-200", 6, date(2014, 8, 20), drive="Hydraulic")
add_contract(maple_c, maple, "Basic", date(2021, 7, 1), date(2025, 6, 30))  # expired
add_contact(maple_c, maple, "Daniel", "Brooks", "HOA Board President", ph="+1-510-555-0177")
inspection(maple_eq, date(2025, 9, 2), "Pass with Notes", "Hydraulic oil level low; recommend service")
for _ in range(9):
    t = NOW - timedelta(minutes=rng.randint(60, 8000))
    add_event(maple_eq, rng.choice(["LEV-221", "LEV-221", "HYD-801"]), t, rng.randint(1, 6))
add_wo(maple_eq, "Repair", "T-502", datetime(2025, 3, 11, 9, 0), component="Hydraulic control valve",
       findings="Valve leak-down causing re-leveling", action="Rebuilt control valve", hours=4.0)

# ---- Alt 3: Lakeside Medical - telemetry shows obstruction, no faults ------
lake_c = add_customer("Lakeside Health Partners", "Healthcare", "Strategic", "Walnut Creek")
lake = add_site(lake_c, "Lakeside Medical Center", "Walnut Creek", "SF Bay - East", 5, "Hospital", 37.9020, -122.0610)
lake_units = [add_equipment(lake, "Elevator %d" % i, "Gen360", "OTC-600", 5, date(2022, 2, 1)) for i in (1, 2, 3, 4)]
add_contract(lake_c, lake, "Comprehensive", date(2023, 1, 1), date(2027, 12, 31))
add_contact(lake_c, lake, "Mary", "Green", "Director of Facilities", ph="+1-925-555-0164")
for u in lake_units:
    pm_history(u, ["T-506", "T-507"])
    inspection(u, date(2026, 5, 10))
# Elevator 2 near the ER: many door reversals, no fault codes
for _ in range(64):
    t = NOW - timedelta(minutes=rng.randint(30, 7 * 24 * 60))
    add_event(lake_units[1], rng.choice(["DPD-305"] * 5 + ["DPD-306"]), t, rng.choice([1, 1, 1, 2]), door="Opening")

# ---- Alt 1: entrapment - Bayview Plaza has a live ENT-900 -----------------
bay_c = add_customer("Bayview Retail Holdings", "Retail", "Key", "San Jose")
bay = add_site(bay_c, "Bayview Plaza", "San Jose", "SF Bay - South", 4, "Retail", 37.3310, -121.8900)
bay_units = [add_equipment(bay, lbl, "Gen2", "OTC-400", 4, date(2017, 6, 1)) for lbl in ("North Car", "South Car")]
add_contract(bay_c, bay, "Standard", date(2025, 1, 1), date(2027, 12, 31))
add_contact(bay_c, bay, "Kevin", "Hall", "Mall Security Supervisor", ph="+1-408-555-0133")
for u in bay_units:
    pm_history(u, ["T-504", "T-508"], months=6)
e1 = add_event(bay_units[0], "SAF-501", NOW - timedelta(minutes=6), 2, "Between Floors", "Up")
add_event(bay_units[0], "ALM-910", NOW - timedelta(minutes=4), 2, "Between Floors", "Idle")
add_event(bay_units[0], "ENT-900", NOW - timedelta(minutes=4, seconds=30), 2, "Between Floors", "Idle", related=e1)

# ================================================================== BACKGROUND
segments = [("Commercial Office", "Office"), ("Residential", "Residential"), ("Hospitality", "Hotel"),
            ("Healthcare", "Clinic"), ("Government", "Civic"), ("Retail", "Retail")]
names = ["Summit Realty Group", "Pacific Crest REIT", "Golden Gate Hotels", "Bay Civic Services", "Redwood Living",
         "Silicon Plaza Partners", "Mission Bay Offices", "Harborview Condos", "Evergreen Senior Living",
         "Northpoint Hospitality", "Peninsula Medical Group", "Union Square Holdings"]
controllers = {"OTC-200": ("Gen2", "Hydraulic"), "OTC-400": ("Gen2", "Geared traction"),
               "OTC-600": ("Gen360", "Gearless traction")}
fault_codes = [c[0] for c in ERROR_CODES if c[3] in ("Low", "Medium", "High") and c[0] not in ("DLC-101", "SEC-960")]
bg_units = []
for nm in names:
    seg, btype = rng.choice(segments)
    region = rng.choice(list(REGION_CITIES))
    city, lat, lon = rng.choice(REGION_CITIES[region])
    cu = add_customer(nm, seg, rng.choice(["Standard", "Standard", "Key"]), city)
    for s in range(rng.randint(1, 3)):
        floors = rng.randint(3, 30)
        st = add_site(cu, "%s %s" % (rng.choice(["One", "Two", "The", "Park", "Bay", "Oak"]),
                                       rng.choice(["Center", "Tower", "Plaza", "Commons", "Court", "Place"])),
                      city, region, floors, btype, lat + rng.uniform(-0.03, 0.03), lon + rng.uniform(-0.03, 0.03))
        tier = rng.choice(["Comprehensive", "Standard", "Standard", "Basic"])
        start = date(rng.randint(2019, 2025), 1, 1)
        add_contract(cu, st, tier, start, start.replace(year=start.year + rng.choice([1, 3, 5])))
        add_contact(cu, st, rng.choice(FIRST), rng.choice(LAST),
                    rng.choice(["Building Manager", "Chief Engineer", "Property Manager"]))
        ctrl = "OTC-200" if floors <= 6 and rng.random() < 0.6 else rng.choice(["OTC-400", "OTC-600"])
        pl, drive = controllers[ctrl]
        install = date(rng.randint(2008, 2025), rng.randint(1, 12), rng.randint(1, 28))
        region_techs = [t[0] for t in TECHS if t[2] == region]
        for u in range(rng.randint(1, 4)):
            eq = add_equipment(st, "Car %s" % "ABCD"[u], pl, ctrl, floors, install, drive=drive,
                               status="Out of Service" if rng.random() < 0.03 else "In Service")
            bg_units.append((eq, cu, st, floors, region_techs))
            pm_history(eq, region_techs, months=rng.randint(4, 12))
            inspection(eq, TODAY - timedelta(days=rng.randint(30, 400)),
                       rng.choice(["Pass"] * 8 + ["Pass with Notes", "Fail"]))

# background telemetry, cases and repairs
for eq, cu, st, floors, techs in bg_units:
    contact = next(c["contact_id"] for c in rows["contacts"] if c["site_id"] == st)
    for _ in range(rng.randint(5, 40)):
        add_event(eq, rng.choice(["DPD-305"] * 4 + ["DPD-306", "COM-701", "PWR-601"]),
                  NOW - timedelta(minutes=rng.randint(60, 90 * 24 * 60)), rng.randint(1, floors), door="Opening")
    for _ in range(rng.randint(0, 3)):
        code = rng.choice(fault_codes)
        cat = next(c for c in ERROR_CODES if c[0] == code)
        t = NOW - timedelta(days=rng.randint(8, 330), hours=rng.randint(0, 10))
        fl = rng.randint(1, floors)
        add_event(eq, code, t, fl, rng.choice(["At Landing", "Between Floors"]))
        opened = t + timedelta(hours=rng.randint(1, 20))
        case = add_case(cu, st, eq, contact, opened, rng.choice(["Phone", "Phone", "Email", "Portal", "IoT Auto"]),
                        rng.choice(["Breakdown", "Door Issue", "Noise", "Leveling"]), cat[2],
                        "Customer reported: " + cat[7], rng.choice(["P2", "P3", "P3"]), "Closed",
                        resolved=opened + timedelta(hours=rng.randint(3, 48)), resolution=cat[6])
        part = rng.choice([p for p in PARTS if cat[1].split()[0] in p[2] or rng.random() < 0.15] or [PARTS[4]])
        replaced = rng.random() < 0.5
        add_wo(eq, "Repair", rng.choice(techs), opened + timedelta(hours=rng.randint(1, 6)), floor=fl,
               component=part[1] if replaced else cat[1], replaced=int(replaced), case_id=case,
               findings=cat[5].split(";")[0], action=cat[6], billable=0,
               parts=[(part[0], 1)] if replaced else ())

# a couple of open background cases so the agent must check for duplicates
for eq, cu, st, floors, techs in rng.sample(bg_units, 4):
    contact = next(c["contact_id"] for c in rows["contacts"] if c["site_id"] == st)
    opened = NOW - timedelta(hours=rng.randint(2, 30))
    case = add_case(cu, st, eq, contact, opened, "Phone", "Noise", "Grinding noise when car starts",
                    "Tenant complaints about noise on start", "P3", "In Progress", sla_hours=8)
    add_wo(eq, "Repair", rng.choice(techs), NOW + timedelta(hours=rng.randint(2, 20)), case_id=case,
           component="Drive / machine", status="Scheduled")

# ------------------------------------------------------------------ reference tables
for c in ERROR_CODES:
    rows["error_code_catalog"].append(dict(zip(
        ["error_code", "category", "description", "severity", "is_safety_related", "likely_causes",
         "recommended_action", "customer_friendly_explanation", "auto_case_threshold_7d"], c)))
for p in PARTS:
    rows["parts"].append(dict(zip(["part_no", "description", "component_category", "unit_price_usd",
                                   "part_warranty_days", "compatible_controllers"], p)))
for t in TECHS:
    rows["technicians"].append(dict(
        technician_id=t[0], full_name=t[1], service_region=t[2], home_base_city=t[3], certifications=t[4],
        skill_level=t[5], phone=phone(), shift_status=t[6], current_lat=t[7], current_lon=t[8]))

# ================================================================== WRITE
ORDER = ["customers", "sites", "equipment", "service_contracts", "parts", "warranties", "contacts",
         "error_code_catalog", "error_events", "technicians", "service_cases", "work_orders",
         "work_order_parts", "inspections", "access_credentials"]

if os.path.exists(DB):
    os.remove(DB)
con = sqlite3.connect(DB)
con.execute("PRAGMA foreign_keys = ON")
con.executescript(open(os.path.join(HERE, "schema.sql")).read())
for t in ORDER:
    data = rows[t]
    if t == "error_events":
        data.sort(key=lambda r: r["event_ts"])
    cols = list(data[0].keys())
    con.executemany("INSERT INTO %s (%s) VALUES (%s)" % (t, ",".join(cols), ",".join("?" * len(cols))),
                    [tuple(r[c] for c in cols) for r in data])
con.commit()

os.makedirs(os.path.join(HERE, "csv"), exist_ok=True)
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
except ImportError:
    wb = None
for t in ORDER:
    cur = con.execute("SELECT * FROM %s" % t)
    cols = [d[0] for d in cur.description]
    data = cur.fetchall()
    with open(os.path.join(HERE, "csv", t + ".csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        w.writerows(data)
    if wb:
        ws = wb.create_sheet(t[:31])
        ws.append(cols)
        for r in data:
            ws.append(list(r))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            width = max([len(str(c))] + [len(str(r[i - 1] or "")) for r in data[:200]])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(width + 2, 10), 60)
    print("%-20s %6d rows" % (t, len(data)))
if wb:
    wb.save(os.path.join(HERE, "otis_hackathon_data.xlsx"))
con.close()
print("Wrote", DB)
